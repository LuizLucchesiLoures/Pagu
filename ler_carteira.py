"""Lê células específicas da planilha Carteira.xlsx e salva um CSV."""
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import yfinance as yf
from openpyxl import load_workbook

BASE = Path(__file__).parent
ARQUIVO = BASE / "data_input" / "Carteira.xlsx"
ABA = "ACOES III"
LINHA_INICIAL, LINHA_FINAL = 5, 20
COLUNAS = {
    "B": "ativo",
    "E": "quantidade",
    "C": "preco_medio_compra",
    "J": "proventos_recebidos",
}
# Nomes da planilha que não são o ticker completo
APELIDOS = {"VALE": "VALE3", "AXIA": "AXIA3", "GRO": "GROP31"}
# Estimativa de dividendos feita no passado ("Legacy"): W = ano, X = valor
LEGACY_LINHAS = (3, 42)
FIGURA_LEGACY = BASE / "fig" / "legacy_dividendos.png"
SAIDA_LEGACY = BASE / "data_output" / "estimativa_legacy.csv"
# Aba "resumo": proventos por tipo e aportes/retiradas por ano
ABA_RESUMO = "resumo"
INTERVALO_RESUMO = "A1:R11"
SAIDA_RESUMO_ANUAL = BASE / "data_output" / "resumo_anual.csv"
SAIDA = BASE / "data_output" / "carteira_lida.csv"
FIGURA_PIZZA = BASE / "fig" / "carteira_pizza.png"
SAIDA_RESUMO = BASE / "data_output" / "resumo_carteira.csv"
FIGURA_RESUMO = BASE / "fig" / "resumo_carteira.png"


def ler_carteira_excel(caminho=ARQUIVO, aba=ABA):
    """Lê só as colunas/linhas definidas e devolve um DataFrame sem linhas vazias."""
    ws = load_workbook(caminho, data_only=True, read_only=True)[aba]
    dados = {nome: [] for nome in COLUNAS.values()}
    for linha in range(LINHA_INICIAL, LINHA_FINAL + 1):
        for col, nome in COLUNAS.items():
            dados[nome].append(ws[f"{col}{linha}"].value)
    df = pd.DataFrame(dados).dropna(subset=["ativo"])
    df["ativo"] = df["ativo"].astype(str).str.strip().str.upper()
    df["ativo"] = df["ativo"].replace(APELIDOS)
    return df.reset_index(drop=True)


def ler_estimativa_legacy(caminho=ARQUIVO, aba=ABA, linhas=LEGACY_LINHAS):
    """Lê ano (coluna W) e dividendos estimados (coluna X); descarta linhas vazias."""
    ws = load_workbook(caminho, data_only=True, read_only=True)[aba]
    registros = [(ws[f"W{i}"].value, ws[f"X{i}"].value) for i in range(linhas[0], linhas[1] + 1)]
    df = pd.DataFrame(registros, columns=["ano", "dividendos_legacy"]).dropna()
    df["ano"] = df["ano"].astype(int)
    return df.reset_index(drop=True)


def plotar_legacy(legacy, modelo=None, realizado=None, caminho_saida=FIGURA_LEGACY, ano_final=None):
    """Barras por ano (R$ mil): Legacy (planilha), modelo e proventos realizados.

    `modelo` e `realizado` são Series indexadas por ano. Se `ano_final` for
    informado, os anos posteriores não aparecem.
    """
    series = [("Legacy (planilha)", legacy.set_index("ano")["dividendos_legacy"], "#2f5597")]
    if modelo is not None:
        series.append(("Modelo bayesiano (mediana)", modelo.set_axis(modelo.index.astype(int)), "#e67e22"))
    if realizado is not None:
        series.append(("Realizado (dividendos + juros)", realizado, "#2e8b57"))
    if ano_final is not None:
        series = [(n, sr[sr.index <= ano_final], c) for n, sr, c in series]
    largura = 0.8 / len(series)
    figura, eixo = plt.subplots(figsize=(13, 6))
    for i, (nome, serie, cor) in enumerate(series):
        deslocamento = (i - (len(series) - 1) / 2) * largura
        eixo.bar(serie.index + deslocamento, serie / 1000, width=largura, color=cor, label=nome)
    eixo.set_xlabel("Ano")
    eixo.set_ylabel("Dividendos / proventos (R$ mil)")
    eixo.set_title("Dividendos por ano: Legacy × modelo × realizado")
    anos = sorted(set().union(*(set(serie.index) for _, serie, _ in series)))
    eixo.set_xticks(anos[::2])
    eixo.tick_params(axis="x", rotation=60)
    eixo.grid(axis="y", alpha=0.3)
    eixo.legend()
    caminho_saida.parent.mkdir(exist_ok=True)
    figura.savefig(caminho_saida, dpi=150, bbox_inches="tight")
    plt.close(figura)


def ler_resumo_anual(caminho=ARQUIVO, aba=ABA_RESUMO, intervalo=INTERVALO_RESUMO):
    """Lê a tabela da aba resumo (itens nas linhas, anos nas colunas).

    A coluna TOTAL vem da planilha. O rótulo "total" aparece duas vezes na
    planilha; aqui ficam "total proventos" (linha 6) e "saldo aportes" (linha 11).
    Células vazias ou só com espaço viram NaN.
    """
    ws = load_workbook(caminho, data_only=True)[aba]
    linhas = [[c.value for c in linha] for linha in ws[intervalo]]
    cabecalho = ["item"] + [str(c).strip() for c in linhas[0][1:]]
    itens = []
    for i, linha in enumerate(linhas[1:], start=2):
        if linha[0] is None:
            continue
        nome = str(linha[0]).strip()
        if nome == "total":
            nome = "total proventos" if i < 8 else "saldo aportes"
        itens.append([nome] + linha[1:])
    df = pd.DataFrame(itens, columns=cabecalho).set_index("item")
    df = df.apply(lambda col: pd.to_numeric(col.map(lambda v: np.nan if isinstance(v, str) else v)))
    return df


def gerar_grafico_legacy(ano_final=None):
    """Lê os CSVs já gerados (Legacy, modelo e realizado) e salva o gráfico de barras."""
    legacy = pd.read_csv(SAIDA_LEGACY, sep=";", decimal=",")
    arquivo_modelo = BASE / "data_output" / "dividendos_carteira_por_ano.csv"
    modelo = None
    if arquivo_modelo.exists():
        modelo = pd.read_csv(arquivo_modelo, sep=";", decimal=",", index_col=0).loc["TOTAL"]
    anual = pd.read_csv(SAIDA_RESUMO_ANUAL, sep=";", decimal=",", index_col=0)
    realizado = anual.loc["dividendos", ["2025", "2026"]] + anual.loc["juros", ["2025", "2026"]]
    realizado.index = realizado.index.astype(int)
    plotar_legacy(legacy, modelo, realizado, ano_final=ano_final)
    print(f"Gráfico Legacy salvo em {FIGURA_LEGACY}")


def buscar_cotacoes_yahoo(tickers):
    """Busca no Yahoo o último fechamento disponível (R$) de cada ticker."""
    cotacoes = {}
    for ticker in tickers:
        try:
            fechamentos = yf.Ticker(ticker).history(period="5d")["Close"].dropna()
            cotacoes[ticker] = float(fechamentos.iloc[-1])
        except Exception as erro:
            print(f"Não foi possível obter a cotação de {ticker}: {erro}")
            cotacoes[ticker] = np.nan
    return pd.Series(cotacoes, dtype=float)


def resumir_carteira(carteira, cotacoes):
    """Resumo por ativo (só posições com quantidade > 0) e linha TOTAL.

    `cotacoes` é indexada por ticker do Yahoo (ativo + ".SA").
    Lucro = valor atual - investido; lucro com dividendos soma os proventos
    já recebidos. Percentuais são sobre o valor investido.
    """
    posicoes = carteira[carteira["quantidade"] > 0].copy()
    posicoes["cotacao"] = (posicoes["ativo"] + ".SA").map(cotacoes)
    posicoes["investido"] = posicoes["quantidade"] * posicoes["preco_medio_compra"]
    posicoes["valor_atual"] = posicoes["quantidade"] * posicoes["cotacao"]
    posicoes = posicoes.rename(columns={"proventos_recebidos": "dividendos"})
    posicoes["dividendos"] = pd.to_numeric(posicoes["dividendos"], errors="coerce")
    resumo = posicoes[
        ["ativo", "quantidade", "preco_medio_compra", "cotacao",
         "investido", "valor_atual", "dividendos"]
    ].set_index("ativo")
    com_cotacao = resumo[resumo["cotacao"].notna()]
    if len(com_cotacao) < len(resumo):
        print("Aviso: ativos sem cotação ficam fora do TOTAL:",
              ", ".join(resumo.index[resumo["cotacao"].isna()]))
    resumo.loc["TOTAL", ["investido", "valor_atual", "dividendos"]] = com_cotacao[
        ["investido", "valor_atual", "dividendos"]
    ].sum()
    resumo["lucro"] = resumo["valor_atual"] - resumo["investido"]
    resumo["lucro_pct"] = resumo["lucro"] / resumo["investido"]
    resumo["lucro_com_dividendos"] = resumo["lucro"] + resumo["dividendos"]
    resumo["lucro_com_dividendos_pct"] = resumo["lucro_com_dividendos"] / resumo["investido"]
    return resumo


def plotar_tabela_resumo(resumo, caminho_saida=FIGURA_RESUMO):
    """Figura com a tabela resumo da carteira (valores em R$)."""
    reais = "{:,.0f}".format
    pct = "{:+.1%}".format
    tabela = pd.DataFrame(
        {
            "Ativo": resumo.index,
            "Qtd.": resumo["quantidade"].map("{:,.0f}".format).replace("nan", ""),
            "Preço médio": resumo["preco_medio_compra"].map("{:,.2f}".format).replace("nan", ""),
            "Cotação": resumo["cotacao"].map("{:,.2f}".format).replace("nan", ""),
            "Investido": resumo["investido"].map(reais),
            "Valor atual": resumo["valor_atual"].map(reais),
            "Lucro": resumo["lucro"].map(reais),
            "Lucro %": resumo["lucro_pct"].map(pct),
            "Dividendos": resumo["dividendos"].map(reais),
            "Lucro+div.": resumo["lucro_com_dividendos"].map(reais),
            "Lucro+div. %": resumo["lucro_com_dividendos_pct"].map(pct),
        }
    ).reset_index(drop=True)
    figura, eixo = plt.subplots(figsize=(15, 0.5 * len(tabela) + 1.5))
    eixo.axis("off")
    desenho = eixo.table(
        cellText=tabela.values, colLabels=tabela.columns, cellLoc="center", loc="center"
    )
    desenho.auto_set_font_size(False)
    desenho.set_fontsize(10)
    desenho.scale(1, 1.5)
    colunas_lucro = [tabela.columns.get_loc(c) for c in ("Lucro", "Lucro %", "Lucro+div.", "Lucro+div. %")]
    valores_lucro = resumo[["lucro", "lucro", "lucro_com_dividendos", "lucro_com_dividendos"]].to_numpy()
    for (linha, coluna), celula in desenho.get_celld().items():
        if linha == 0:
            celula.set_facecolor("#d9e2f3")
            celula.set_text_props(weight="bold")
            continue
        if linha == len(tabela):
            celula.set_text_props(weight="bold")
            celula.set_facecolor("#f2f2f2")
        if coluna in colunas_lucro:
            valor = valores_lucro[linha - 1][colunas_lucro.index(coluna)]
            celula.set_text_props(color="#1a7f37" if valor >= 0 else "#c0392b")
    figura.suptitle("Resumo da carteira (R$)", fontsize=13)
    caminho_saida.parent.mkdir(exist_ok=True)
    figura.savefig(caminho_saida, dpi=150, bbox_inches="tight")
    plt.close(figura)


def plotar_pizza_carteira(carteira, caminho_saida=FIGURA_PIZZA, ativos=None):
    """Pizza com a composição da carteira (quantidade × preço médio por ativo).

    Se `ativos` for informado (tickers com ou sem ".SA"), só eles entram.
    """
    if ativos is not None:
        ativos = [a.replace(".SA", "") for a in ativos]
        carteira = carteira[carteira["ativo"].isin(ativos)]
    investido = (carteira["quantidade"] * carteira["preco_medio_compra"]).set_axis(carteira["ativo"])
    investido = investido[investido > 0].sort_values(ascending=False)
    figura, eixo = plt.subplots(figsize=(9, 9))
    eixo.pie(
        investido,
        labels=investido.index,
        autopct="%1.1f%%",
        colors=plt.cm.tab20.colors[: len(investido)],
        startangle=90,
        counterclock=False,
        pctdistance=0.8,
        wedgeprops={"edgecolor": "white"},
    )
    eixo.set_title(f"Composição da carteira (valor investido: R$ {investido.sum():,.0f})")
    caminho_saida.parent.mkdir(exist_ok=True)
    figura.savefig(caminho_saida, dpi=150, bbox_inches="tight")
    plt.close(figura)


if __name__ == "__main__":
    carteira = ler_carteira_excel()
    SAIDA.parent.mkdir(exist_ok=True)
    carteira.to_csv(SAIDA, index=False, decimal=",", sep=";")
    print(carteira.to_string())
    print(f"\nSalvo em {SAIDA}")
    resumo_anual = ler_resumo_anual()
    resumo_anual.to_csv(SAIDA_RESUMO_ANUAL, sep=";", decimal=",", float_format="%.2f")
    print(f"Resumo anual (aba resumo) salvo em {SAIDA_RESUMO_ANUAL}")
    legacy = ler_estimativa_legacy()
    legacy.to_csv(SAIDA_LEGACY, index=False, sep=";", decimal=",", float_format="%.2f")
    print(f"Legacy: {len(legacy)} anos ({legacy['ano'].min()}-{legacy['ano'].max()}) em {SAIDA_LEGACY}")
    gerar_grafico_legacy()
    plotar_pizza_carteira(carteira)
    print(f"Pizza salva em {FIGURA_PIZZA}")
    cotacoes = buscar_cotacoes_yahoo((carteira.loc[carteira["quantidade"] > 0, "ativo"] + ".SA"))
    resumo = resumir_carteira(carteira, cotacoes)
    resumo.to_csv(SAIDA_RESUMO, sep=";", decimal=",", float_format="%.4f")
    plotar_tabela_resumo(resumo)
    print(resumo.round(2).to_string())
    print(f"Resumo salvo em {SAIDA_RESUMO} e {FIGURA_RESUMO}")
